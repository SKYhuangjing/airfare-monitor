"""Ctrip complete round-trip combinations, with no browser or network access."""
from __future__ import annotations

import copy
import json
import tempfile
import unittest
from dataclasses import replace
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from functools import lru_cache
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock
from urllib.parse import parse_qs, urlsplit

from openpyxl import load_workbook

from airfare_monitor.collector import QunarBrowserSession
from airfare_monitor.config import BrowserSettings, load_settings
from airfare_monitor.ctrip_parser import parse_ctrip_payload
from airfare_monitor.ctrip_search import CtripSearchCapture, SEARCH_PREFIX, build_ctrip_search_url
from airfare_monitor.errors import CollectionError, IncompleteResponseError, ParseError
from airfare_monitor.models import EtdWindow, LegConfig, LegResult, LegStatus, PreferredSchedule, RunReport, RunStatus
from airfare_monitor.service import MonitorService

ROOT = Path(__file__).resolve().parents[1]
NOW = datetime(2026, 10, 12, 15)
BATCH = "https://flights.ctrip.com" + SEARCH_PREFIX + "batchSearch"
PULL = "https://flights.ctrip.com" + SEARCH_PREFIX + "pull/search-round"


def leg(**changes):
    return replace(LegConfig("ctrip-round", True, "PVG", "KUL", date(2026, 10, 14),
                             EtdWindow(time(0), time(23, 59)), False, Decimal("2000"), 200, 1, 0, "economy",
                             return_date=date(2026, 10, 17), return_etd_window=EtdWindow(time(0), time(23, 59))), **changes)


@lru_cache(maxsize=1)
def recorded():
    return json.loads((Path(__file__).parent / "fixtures/ctrip_roundtrip_final_sanitized.json").read_text(encoding="utf-8"))


def payload(rows=None, *, finished=True):
    value = copy.deepcopy(recorded()) if rows is None else {"status": 0, "data": {"flightItineraryList": copy.deepcopy(rows)}}
    value["data"]["context"] = {"finished": finished, "flag": 0, "showAuthCode": False, "searchId": "search-round"}
    return value


def direct_row():
    for row in recorded()["data"]["flightItineraryList"]:
        ss = row["flightSegments"]
        if (all(len(s["flightList"]) == 1 and s["stopCount"] == 0 for s in ss)
                and ss[0]["flightList"][0]["departureAirportCode"] == "PVG"
                and ss[1]["flightList"][0]["arrivalAirportCode"] == "PVG"):
            value = copy.deepcopy(row)
            value["priceList"] = [{"adultPrice": 1000, "adultTax": 200, "eligibility": "NOR", "cabin": "Y"}]
            return value
    raise AssertionError("Recorded direct round-trip sample missing")


def request(**changes):
    value = {"transactionID": "round-transaction", "adultCount": 1, "childCount": 0, "infantCount": 0,
             "flightWay": "D", "scope": "i", "segmentNo": 0, "cabin": "Y_S", "flightSegments": [
                 {"departureCityCode": "SHA", "arrivalCityCode": "KUL", "departureDate": "2026-10-14"},
                 {"departureCityCode": "KUL", "arrivalCityCode": "SHA", "departureDate": "2026-10-17"}]}
    value.update(changes)
    return value


def criteria():
    return {"transactionID": "round-transaction", **request()["flightSegments"][0],
            "returnSegment": request()["flightSegments"][1]}


class CtripRoundTripParserTests(unittest.TestCase):
    def test_recorded_combinations_total_and_candidate_limit(self):
        flights, _, observed, eligible = parse_ctrip_payload(recorded(), leg(), NOW)
        self.assertEqual((observed, eligible, len(flights)), (1292, 511, 200))
        self.assertEqual(flights[0].total_price_cny, Decimal("1903"))
        self.assertEqual((flights[0].base_price_cny, flights[0].tax_cny), (Decimal("510"), Decimal("1393")))
        self.assertTrue(all(f.return_itinerary and f.return_itinerary.destination_airport_iata == "PVG" for f in flights))

    def test_city_aggregation_roundtrip_accepts_physical_airports(self):
        flights, _, observed, eligible = parse_ctrip_payload(
            recorded(),
            leg(origin_airport_iata="SHA", origin_airports=("PVG", "SHA")),
            NOW,
        )
        self.assertEqual((observed, eligible, len(flights)), (1292, 1221, 200))

    def test_same_outbound_different_return_remains_distinct_and_price_counts_once(self):
        first = direct_row()
        second = copy.deepcopy(first)
        second["flightSegments"][1]["flightList"][0]["flightNo"] = "MU9999"
        second["flightSegments"][1]["flightList"][0]["marketAirlineCode"] = "MU"
        duplicate = copy.deepcopy(first)
        duplicate["priceList"][0]["adultPrice"] = 2000
        flights, _, _, eligible = parse_ctrip_payload(payload([first, second, duplicate]), leg(), NOW)
        self.assertEqual(eligible, 2)
        self.assertEqual(len({f.flight_signature for f in flights}), 2)
        self.assertEqual(len({f.flight_codes for f in flights}), 1)
        self.assertTrue(all(f.total_price_cny == Decimal("1200") for f in flights))

    def test_return_airport_date_and_time_window_are_enforced(self):
        for key, value in [("arrivalAirportCode", "SHA"), ("departureAirportCode", "SIN"),
                           ("departureDateTime", "2026-10-18 10:00:00")]:
            row = direct_row()
            row["flightSegments"][1]["flightList"][0][key] = value
            with self.subTest(key=key):
                self.assertEqual(parse_ctrip_payload(payload([row]), leg(), NOW)[3], 0)
        row = direct_row()
        self.assertEqual(parse_ctrip_payload(payload([row]), leg(return_etd_window=EtdWindow(time(0), time(0))), NOW)[3], 0)
        self.assertEqual(parse_ctrip_payload(payload([row]), leg(etd_window=EtdWindow(time(0), time(0))), NOW)[3], 0)

    def test_each_direction_applies_its_own_direct_and_layover_rule(self):
        flights, _, _, _ = parse_ctrip_payload(recorded(), leg(direct_only=True, return_direct_only=False, return_max_layover_minutes=240), NOW)
        self.assertTrue(flights)
        self.assertTrue(all(f.is_direct for f in flights))
        self.assertTrue(any(not f.return_itinerary.is_direct for f in flights))
        self.assertTrue(all(f.return_itinerary.layover_minutes is None or f.return_itinerary.layover_minutes <= 240 for f in flights))
        flights, _, _, _ = parse_ctrip_payload(recorded(), leg(direct_only=False, max_layover_minutes=240, return_direct_only=True), NOW)
        self.assertTrue(flights)
        self.assertTrue(any(not f.is_direct for f in flights))
        self.assertTrue(all(f.return_itinerary.is_direct for f in flights))
        self.assertTrue(all(f.layover_minutes is None or f.layover_minutes <= 240 for f in flights))

    def test_direction_order_is_resolved_by_segment_number_and_partial_shape_rejected(self):
        row = direct_row()
        row["flightSegments"].reverse()
        self.assertEqual(parse_ctrip_payload(payload([row]), leg(), NOW)[3], 1)
        for segments in [row["flightSegments"][:1], [row["flightSegments"][0]] * 2]:
            broken = copy.deepcopy(row)
            broken["flightSegments"] = segments
            with self.subTest(segments=len(segments)), self.assertRaises(ParseError):
                parse_ctrip_payload(payload([broken]), leg(), NOW)
        with self.assertRaises(IncompleteResponseError):
            parse_ctrip_payload(payload([row], finished=False), leg(), NOW)

    def test_mixed_cabins_and_group_total(self):
        row = direct_row()
        base = row["priceList"][0]
        row["priceList"] = [dict(base, adultPrice=1, cabin="Y|C"), dict(base, adultPrice=100, adultTax=50, childPrice=80, childTax=20)]
        flights, _, _, _ = parse_ctrip_payload(payload([row]), leg(adult_count=2, child_count=1), NOW)
        self.assertEqual(flights[0].total_price_cny, Decimal("400"))

    def test_preferred_match_keeps_the_return_and_no_secret_metadata(self):
        row = direct_row()
        first = row["flightSegments"][0]["flightList"][0]
        preferred = PreferredSchedule("focus", datetime.fromisoformat(first["departureDateTime"]).time(),
                                      datetime.fromisoformat(first["arrivalDateTime"]).time(), 0, 0, 0, "PVG", "KUL")
        row["priceList"][0]["routeSearchToken"] = "PRIVATE_SENTINEL"
        _, matches, _, _ = parse_ctrip_payload(payload([row]), leg(top_n=1, preferred_schedules=(preferred,)), NOW)
        self.assertIsNotNone(matches[0].return_itinerary)
        self.assertNotIn("PRIVATE_SENTINEL", json.dumps(matches[0].raw_item))


class CtripRoundTripCaptureTests(unittest.TestCase):
    def test_url_uses_native_combination_mode_and_both_dates(self):
        url = urlsplit(build_ctrip_search_url(leg(adult_count=2, child_count=1)))
        self.assertEqual(url.path, "/online/list/round-sha-kul")
        self.assertEqual(parse_qs(url.query), {"depdate": ["2026-10-14_2026-10-17"], "isbuildup": ["1"],
                                              "adult": ["2"], "child": ["1"], "infant": ["0"], "cabin": ["y_s"]})

    def test_split_stages_wrong_return_and_old_search_cannot_complete_roundtrip(self):
        for bad in [request(segmentNo=1), request(segmentNo=2), request(segmentNo=False), request(flightWay="S"),
                    request(transactionID="old"), request(flightSegments=request()["flightSegments"][:1])]:
            with self.subTest(request=bad):
                self.assertIsNone(CtripSearchCapture(leg(), "round-transaction").accept(BATCH, bad, payload([direct_row()])))
        for key, value in [("departureDate", "2026-10-18"), ("departureCityCode", "SIN"), ("arrivalCityCode", "BJS")]:
            bad = request()
            bad["flightSegments"][1][key] = value
            with self.subTest(key=key):
                self.assertIsNone(CtripSearchCapture(leg(), "round-transaction").accept(BATCH, bad, payload([direct_row()])))

    def test_final_replaces_partial_combinations_and_missing_body_fails(self):
        capture = CtripSearchCapture(leg(), "round-transaction")
        self.assertIsNone(capture.accept(BATCH, request(), payload([direct_row(), direct_row()], finished=False)))
        final = capture.accept(PULL, request(), payload([direct_row()]))
        self.assertEqual(len(final["data"]["flightItineraryList"]), 1)
        with self.assertRaises(ParseError):
            capture.accept(PULL, request(), None)


class CtripRoundTripPipelineTests(unittest.TestCase):
    def session(self, tmp):
        session = QunarBrowserSession(BrowserSettings(False, Path(tmp) / "data/browser-profile", 9438, 45, 75, 2, "https://example.test"))
        session.start = Mock()
        session.tab = Mock()
        session.tab.run_js.return_value = criteria()
        return session

    def test_collector_configures_large_buffer_before_navigation(self):
        with tempfile.TemporaryDirectory() as tmp:
            session = self.session(tmp)
            session.tab.listen.wait.return_value = SimpleNamespace(url=BATCH, request=SimpleNamespace(postData=json.dumps(request())),
                                                                  response=SimpleNamespace(body=payload(), status=200))
            result = session.collect(leg(top_n=3), now=lambda: NOW)
            self.assertEqual((result.completed_response, result.observed_count, result.eligible_count), (True, 1292, 511))
            self.assertEqual((len(result.flights), len(result.candidate_flights)), (3, 200))
            self.assertTrue(result.raw_response["round_trip"])
            calls = [call[0] for call in session.tab.mock_calls]
            self.assertLess(calls.index("listen._driver.run"), calls.index("get"))
            session.tab.listen._driver.run.assert_called_once_with("Network.enable", maxTotalBufferSize=134217728, maxResourceBufferSize=67108864)

    def test_buffer_configuration_or_page_return_mismatch_fails_cleanly(self):
        with tempfile.TemporaryDirectory() as tmp:
            for driver in [None, SimpleNamespace(run=lambda *a, **k: {"error": "unsupported"}),
                           SimpleNamespace(run=Mock(side_effect=RuntimeError("driver stopped")))]:
                session = self.session(tmp)
                session.tab.listen._driver = driver
                with self.assertRaises(CollectionError):
                    session.collect(leg(), now=lambda: NOW)
                session.tab.get.assert_not_called()
                session.tab.listen.stop.assert_called_once()
            session = self.session(tmp)
            wrong = criteria()
            wrong["returnSegment"]["departureDate"] = "2026-10-18"
            session.tab.run_js.return_value = wrong
            with self.assertRaises(CollectionError):
                session.collect(leg(), now=lambda: NOW)
            session.tab.listen.wait.assert_not_called()

    def test_monitor_secondary_confirmation_storage_and_excel_keep_return(self):
        flights, matches, observed, eligible = parse_ctrip_payload(recorded(), leg(), NOW)
        with tempfile.TemporaryDirectory() as tmp:
            settings = load_settings(ROOT / "resources/settings.default.yaml", project_root=tmp)
            settings = replace(settings, collection=replace(settings.collection, secondary_confirmation_delay_seconds=0))
            browser = Mock()
            browser.collect.side_effect = lambda route, **_: LegResult(route, LegStatus.SUCCESS, NOW, flights=flights[:10],
                candidate_flights=flights, preferred_matches=matches, completed_response=True, observed_count=observed, eligible_count=eligible)
            service = MonitorService([leg(top_n=10)], settings, browser=browser, now=lambda: NOW)
            report, output = service.run_once(send_email=False)
            self.assertEqual(report.threshold_confirmed_leg_ids, {leg().id})
            self.assertEqual(browser.collect.call_count, 2)
            stored = service.store.latest_flight_candidates(leg().id)
            self.assertEqual(stored["collection_provider"], "ctrip")
            self.assertEqual(stored["stored_count"], 200)
            self.assertIsNotNone(stored["flights"][0]["return_itinerary"])
            self.assertEqual(service.store.flight_price_reference(leg().id, flights[0].flight_signature, provider="ctrip").first_total_price_cny, Decimal("1903"))
            workbook = load_workbook(output, read_only=True)
            try:
                summary = next(workbook["本次汇总"].iter_rows(min_row=2, max_row=2, values_only=True))
                row = next(workbook["航程1"].iter_rows(min_row=2, max_row=2, values_only=True))
                self.assertEqual(summary[1], "往返")
                self.assertEqual(summary[6].date(), date(2026, 10, 17))
                self.assertEqual(row[23], "/".join(flights[0].return_itinerary.flight_codes))
                self.assertEqual(row[25:27], ("KUL", "PVG"))
            finally:
                workbook.close()

    def test_monitor_history_excludes_oneway_and_other_return_dates(self):
        flight = parse_ctrip_payload(payload([direct_row()]), leg(), NOW)[0][0]
        preferred = PreferredSchedule("focus", flight.etd_local.time(), flight.eta_local.time(), 0, 0, 0, "PVG", "KUL")
        current = leg(preferred_schedules=(preferred,), expected_total_price_cny=None)
        with tempfile.TemporaryDirectory() as tmp:
            settings = load_settings(ROOT / "resources/settings.default.yaml", project_root=tmp)
            browser = Mock()
            browser.collect.side_effect = lambda route, **_: LegResult(route, LegStatus.SUCCESS, NOW, flights=[flight],
                preferred_matches=[flight], completed_response=True)
            service = MonitorService([current], settings, browser=browser, now=lambda: NOW)
            service.store.initialize()
            for index, old_leg in enumerate([
                replace(current, return_date=None, return_etd_window=None),
                replace(current, return_date=date(2026, 10, 18)),
                replace(current, origin_airport_iata="SHA"),
                replace(current, departure_date=date(2026, 10, 13)),
            ]):
                captured = NOW - timedelta(minutes=index + 1)
                old_flight = replace(flight, total_price_cny=Decimal("100"), captured_at=captured)
                old = LegResult(old_leg, LegStatus.SUCCESS, captured, flights=[old_flight],
                                preferred_matches=[old_flight], completed_response=True)
                service.store.save_report(RunReport(f"old-{index}", captured, captured, RunStatus.SUCCESS, [old]))
            result = service.run_once(send_email=False)[0].legs[0]
            self.assertIsNone(result.previous_min_total_cny)
            self.assertIsNone(result.preferred_price_references[0].previous_total_price_cny)
            self.assertEqual(result.preferred_price_references[0].first_total_price_cny, Decimal("1200"))
            self.assertIsNone(result.flight_price_references[flight.flight_signature].previous_total_price_cny)
            # A subsequent query of this exact trip still has a valid baseline.
            later = NOW + timedelta(minutes=1)
            browser.collect.side_effect = lambda route, **_: replace(result, captured_at=later)
            service.now = lambda: later
            again = service.run_once(send_email=False)[0].legs[0]
            self.assertEqual(again.previous_min_total_cny, Decimal("1200"))
            self.assertEqual(again.preferred_price_references[0].previous_total_price_cny, Decimal("1200"))


if __name__ == "__main__":
    unittest.main()
